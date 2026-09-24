#!/usr/bin/env python3
"""
Import IPAM_Export_*.xlsx (export multi-sites produit par l'appli elle-même,
bouton "Export Excel") -> Redis  (IPAM SIW v2)
================================================================================
Format attendu : un onglet par site, colonnes "VLAN ID", "Réseau" (CIDR,
ex. 218.16.156.128/25), "Adresse IP", "Hostname", "Statut", et optionnellement
"Gateway" (selon les cases cochées lors de l'export).

Réutilise import_site() de import_redis.py (même logique de création
site/VLAN/IP, même dédoublonnage par IP).

Usage :
  python3 import_redis_export.py --xlsx IPAM_Export_2026-09-24.xlsx
                                  [--host 127.0.0.1] [--port 6379]
                                  [--password PWD] [--dry-run] [--site NOM]
"""

import sys
import re
import argparse
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:
    sys.exit("Dépendance manquante : pip install openpyxl")

try:
    import redis as _redis_lib
except ImportError:
    sys.exit("Dépendance manquante : pip install redis")

from import_redis import import_site  # réutilise la logique de création Redis

NOW = datetime.now(timezone.utc).isoformat()


def _cidr_to_dotted_mask(bits_str):
    try:
        bits = int(bits_str)
    except (TypeError, ValueError):
        return ''
    if bits < 0 or bits > 32:
        return ''
    mask_int = (0xFFFFFFFF << (32 - bits)) & 0xFFFFFFFF if bits > 0 else 0
    return '.'.join(str((mask_int >> (8 * (3 - i))) & 0xFF) for i in range(4))


def _normalize_status(cell):
    s = str(cell).strip() if cell else ''
    low = s.lower()
    if 'serv' in low:  # "Réservée" / "Reservee"
        return 'Réservée'
    if 'utilis' in low:
        return 'Utilisé'
    return 'Libre'


def _find_col(header, *names):
    names_low = [n.lower() for n in names]
    for i, h in enumerate(header):
        if h and str(h).strip().lower() in names_low:
            return i
    return None


def parse_sheet(ws):
    """Retourne { vlan_id_str: {'network','mask','gateway','ips':[{'ip','hostname','status'}]} }"""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}

    header = rows[0]
    c_vlan = _find_col(header, 'VLAN ID', 'VLAN')
    c_net  = _find_col(header, 'Réseau', 'Reseau', 'Network')
    c_ip   = _find_col(header, 'Adresse IP', 'IP')
    c_host = _find_col(header, 'Hostname')
    c_stat = _find_col(header, 'Statut', 'Status')
    c_gw   = _find_col(header, 'Gateway', 'Passerelle')

    if c_vlan is None or c_ip is None:
        return {}

    vlans = {}
    for row in rows[1:]:
        if len(row) <= max(c_vlan, c_ip):
            continue
        vid_raw = row[c_vlan]
        ip_raw  = row[c_ip]
        if not vid_raw or not ip_raw:
            continue
        ip_str = str(ip_raw).strip()
        if not re.match(r'^\d{1,3}(\.\d{1,3}){3}$', ip_str):
            continue
        vid = str(vid_raw).strip()

        net_raw = str(row[c_net]).strip() if c_net is not None and len(row) > c_net and row[c_net] else ''
        network, mask = '', ''
        if '/' in net_raw:
            network, bits = net_raw.split('/', 1)
            network = network.strip()
            mask = _cidr_to_dotted_mask(bits.strip())
        else:
            network = net_raw

        gateway = str(row[c_gw]).strip() if c_gw is not None and len(row) > c_gw and row[c_gw] else ''
        hostname = str(row[c_host]).strip() if c_host is not None and len(row) > c_host and row[c_host] else ''
        status = _normalize_status(row[c_stat] if c_stat is not None and len(row) > c_stat else None)

        if vid not in vlans:
            vlans[vid] = {'network': network, 'mask': mask, 'gateway': gateway, 'ips': []}
        if not vlans[vid]['gateway'] and gateway:
            vlans[vid]['gateway'] = gateway
        if not vlans[vid]['network'] and network:
            vlans[vid]['network'] = network
            vlans[vid]['mask'] = mask

        vlans[vid]['ips'].append({'ip': ip_str, 'hostname': hostname, 'status': status})

    return vlans


def main():
    parser = argparse.ArgumentParser(description='Importe un export IPAM (Export Excel) dans Redis')
    parser.add_argument('--xlsx', required=True, help='Chemin du fichier Excel')
    parser.add_argument('--host', default='127.0.0.1')
    parser.add_argument('--port', type=int, default=6379)
    parser.add_argument('--password', default=None)
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--site', default=None, help='Importer seulement ce site (nom de feuille, partiel)')
    args = parser.parse_args()

    print(f"\nOuverture de {args.xlsx} ...")
    try:
        wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    except FileNotFoundError:
        sys.exit(f"Fichier introuvable : {args.xlsx}")

    print(f"   {len(wb.sheetnames)} feuilles\n")

    if not args.dry_run:
        try:
            r = _redis_lib.Redis(host=args.host, port=args.port, password=args.password, decode_responses=True)
            r.ping()
            print(f"Redis connecte : {args.host}:{args.port}\n")
        except Exception as e:
            sys.exit(f"Impossible de se connecter a Redis : {e}")
    else:
        r = None
        print("Mode DRY-RUN -- aucune ecriture Redis\n")

    grand_vlans = 0
    grand_ips = 0

    for sheet_name in wb.sheetnames:
        if args.site and args.site.lower() not in sheet_name.lower():
            continue

        site_name = sheet_name.strip().upper()
        print(f"-- {site_name}")

        ws = wb[sheet_name]
        vlans_data = parse_sheet(ws)

        if not vlans_data:
            print("   (aucune donnee IP)\n")
            continue

        n = import_site(r, site_name, vlans_data, dry_run=args.dry_run)

        if not args.dry_run:
            n_vlans = len(vlans_data)
            grand_vlans += n_vlans
            grand_ips += n
            for vid, vd in vlans_data.items():
                net = vd.get('network') or '-'
                print(f"   VLAN {vid:>5}  {net:<22}  {len(vd['ips'])} IPs")
            print(f"   -> {n_vlans} VLANs, {n} IPs inserees\n")
        else:
            grand_ips += n
            grand_vlans += len(vlans_data)
            print()

    print("=" * 55)
    print(f"TOTAL : {grand_vlans} VLANs  |  {grand_ips} IPs")
    if args.dry_run:
        print("(DRY-RUN -- rien n'a ete ecrit)")
    print()


if __name__ == '__main__':
    main()
